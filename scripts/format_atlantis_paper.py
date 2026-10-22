"""Convert the retained TopoFlow manuscript using Atlantis Press's Word template.

Run with the managed document Python runtime. The source paper and downloaded
reference are read-only inputs. No publisher macros are executed.
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import json
import re
import urllib.request

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from lxml import etree as E

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / '.paper_revision/atlantis_format'
SOURCE = ROOT / 'TopoFlow_Physics_Guided_Graph_Attention_ICISEA2026.docx'
REFERENCE = WORK / 'template/splnproc2311_reference.docx'
OUTPUT = ROOT / 'TopoFlow_Atlantis_ICISEA2026.docx'
TEMPLATE_URL = 'https://assets.ctfassets.net/b4d0t67iwrjs/61D4gQKQaJuqyRIB8i5h8h/0e1b497360ae6c6d0e04fdf5d776a9e2/Word_Template.zip'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()

def reference_content_digest(path):
    digest = hashlib.sha256()
    with ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            digest.update(name.encode() + b'\0' + hashlib.sha256(archive.read(name)).digest())
    return digest.hexdigest()

def prepare_reference():
    """Recover the exact publisher input without executing the supplied macros."""
    if REFERENCE.is_file():
        return
    REFERENCE.parent.mkdir(parents=True, exist_ok=True)
    package = REFERENCE.parent / 'Word_Template.zip'
    if not package.is_file():
        package.write_bytes(urllib.request.urlopen(TEMPLATE_URL, timeout=30).read())
    assert sha(package) == 'db03865ce8a5c654c4deb8bdaacb8b6ff76f37fe3391e9c32e4815d7edcf8c24'
    with ZipFile(package) as archive:
        macro_doc = archive.read('splnproc2311.docm')
    assert hashlib.sha256(macro_doc).hexdigest() == '1b7358d3f5f783e8fb2a0db194837ff637fda7aadeb7c27e953dac3df9725821'
    with ZipFile(BytesIO(macro_doc)) as source_zip, ZipFile(REFERENCE, 'w', ZIP_DEFLATED) as target_zip:
        for name in source_zip.namelist():
            if any(token in name.lower() for token in ['vbaproject', 'vbadata', 'customui/']):
                continue
            data = source_zip.read(name)
            if name == '[Content_Types].xml' or name.endswith('.rels'):
                root = E.fromstring(data)
                for element in list(root):
                    if 'macroEnabled.main' in element.get('ContentType', ''):
                        element.set('ContentType', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml')
                    attributes = ' '.join(element.attrib.values()).lower()
                    if any(token in attributes for token in ['vba', 'customui']):
                        root.remove(element)
                data = E.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
            target_zip.writestr(name, data)

prepare_reference()
assert sha(SOURCE) == 'd9b7f24f5f74370cf1491e4172ec5a9e750cba27643fb9a6c4b1002c4ce281cf'
assert reference_content_digest(REFERENCE) == 'a50ccb0a5614fa967fabb9cd20a6eced773abfe26b8ddf09340d7b957e14f8b1'
assert (WORK / 'artifact.md').is_file() or (ROOT / 'docs/atlantis_format_conversion.md').is_file()

source = Document(SOURCE)
doc = Document(REFERENCE)
original = list(source._element.body)
patterns = [deepcopy(e) for e in doc._element.body]
assert len(original) == 132
section_xml = E.tostring(doc._element.body.sectPr)
for element in list(doc._element.body):
    if element.tag != qn('w:sectPr'):
        doc._element.body.remove(element)

# Preserve every template style except the explicit Normal size and Title alias.
doc.styles['Normal'].font.size = Pt(10)
title_style = doc.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)
title_style.base_style = doc.styles['papertitle']
for element in list(title_style.element):
    if element.tag == qn('w:rPr'):
        title_style.element.remove(element)

def append(element):
    doc._element.body.insert(len(doc._element.body) - 1, element)
    return Paragraph(element, doc._body) if element.tag == qn('w:p') else element

def paragraph(text='', style='Normal', bold_prefix=None):
    role_index = {'Title': 0, 'author': 1, 'address': 3, 'abstract': 4,
                  'keywords': 5, 'heading1': 6, 'heading2': 7, 'p1a': 8,
                  'Normal': 9, 'equation': 15, 'figurecaption': 18,
                  'tablecaption': 12, 'referenceitem': 23, 'acknowlegments': 21}
    element = deepcopy(patterns[role_index.get(style, 9)])
    for child in list(element):
        if child.tag != qn('w:pPr'):
            element.remove(child)
    p = append(element)
    p.style = doc.styles[style]
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p

def no_number(p):
    pr = p._p.get_or_add_pPr()
    for n in pr.findall(qn('w:numPr')):
        pr.remove(n)
    num = OxmlElement('w:numPr')
    ident = OxmlElement('w:numId'); ident.set(qn('w:val'), '0')
    num.append(ident); pr.append(num)

def heading(text, level=1):
    previous = list(doc._element.body)[-2]
    p = paragraph(text, 'heading' + str(level))
    # The template suppresses space only between a section and its first subsection.
    if level == 2 and previous.xpath('./w:pPr/w:pStyle/@w:val') != ['heading1']:
        p.paragraph_format.space_before = None
    p.paragraph_format.keep_with_next = True
    return p

def copy_math(element):
    copied = deepcopy(element)
    pr = copied.find(qn('w:pPr'))
    if pr is not None:
        copied.remove(pr)
    p = append(copied)
    p.style = doc.styles['equation']
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True

def copy_figure(element):
    copied = deepcopy(element)
    pr = copied.find(qn('w:pPr'))
    if pr is not None:
        copied.remove(pr)
    for embed in copied.xpath('.//*[@r:embed]'):
        old_rid = embed.get(qn('r:embed'))
        blob = source.part.related_parts[old_rid].blob
        rid, _ = doc.part.get_or_add_image(BytesIO(blob))
        embed.set(qn('r:embed'), rid)
    # Enlarge the same images proportionately within the publisher's text area.
    for extent in copied.xpath('.//wp:extent'):
        old_w, old_h = int(extent.get('cx')), int(extent.get('cy'))
        is_workflow = 'workflow' in ''.join(copied.xpath('.//wp:docPr/@descr')).lower()
        new_w = int(Inches(4.4 if is_workflow else 4.0))
        new_h = round(old_h * new_w / old_w)
        extent.set('cx', str(new_w)); extent.set('cy', str(new_h))
        for ext in copied.xpath('.//a:xfrm/a:ext'):
            ext.set('cx', str(new_w)); ext.set('cy', str(new_h))
    p = append(copied)
    p.style = doc.styles['image']
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True

def copy_table(element):
    table = deepcopy(element)
    widths = [int(e.get(qn('w:w'))) for e in table.find(qn('w:tblGrid'))]
    target_width = 6918
    new_widths = [round(w * target_width / sum(widths)) for w in widths]
    new_widths[-1] += target_width - sum(new_widths)
    for col, width in zip(table.find(qn('w:tblGrid')), new_widths):
        col.set(qn('w:w'), str(width))
    props = table.find(qn('w:tblPr'))
    for n in list(props):
        if n.tag in {qn('w:tblBorders'), qn('w:tblCellMar')}:
            props.remove(n)
    width = props.find(qn('w:tblW'))
    width.set(qn('w:w'), str(target_width)); width.set(qn('w:type'), 'dxa')
    margins = OxmlElement('w:tblCellMar')
    for side, amount in [('left', 80), ('right', 80), ('top', 45), ('bottom', 45)]:
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(amount)); e.set(qn('w:type'), 'dxa')
        margins.append(e)
    props.append(margins)
    rows = table.findall(qn('w:tr'))
    for row_index, row in enumerate(rows):
        trpr = row.find(qn('w:trPr'))
        if trpr is None:
            trpr = OxmlElement('w:trPr'); row.insert(0, trpr)
        for n in trpr.findall(qn('w:trHeight')):
            trpr.remove(n)
        if row_index == 0 and trpr.find(qn('w:tblHeader')) is None:
            trpr.append(OxmlElement('w:tblHeader'))
        for cell, width in zip(row.findall(qn('w:tc')), new_widths):
            cp = cell.find(qn('w:tcPr'))
            cp.find(qn('w:tcW')).set(qn('w:w'), str(width))
            for old in cp.findall(qn('w:tcBorders')):
                cp.remove(old)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'bottom', 'left', 'right']:
                border = OxmlElement('w:' + side)
                visible = (row_index == 0 and side in ['top', 'bottom']) or (row_index == len(rows)-1 and side == 'bottom')
                border.set(qn('w:val'), 'single' if visible else 'nil')
                if visible:
                    border.set(qn('w:sz'), '8' if side == 'top' or row_index == len(rows)-1 else '4')
                    border.set(qn('w:color'), '000000')
                borders.append(border)
            cp.append(borders)
            for pe in cell.findall(qn('w:p')):
                p = Paragraph(pe, doc._body)
                p.style = doc.styles['p1a']
                no_number(p)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.keep_with_next = row_index < len(rows)-1
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if row_index else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = 'Times New Roman'; run.font.size = Pt(9)
                    run.bold = row_index == 0
    append(table)

text = lambda e: ''.join(e.xpath('.//w:t/text()'))
title = text(original[0])
paragraph(title, 'Title')
names = [cell.text.split('\n')[0] for cell in source.tables[0].rows[0].cells]
emails = [cell.text.split('\n')[1] for cell in source.tables[0].rows[0].cells]
authors = paragraph('', 'author')
authors.add_run(names[0]); authors.add_run('*').font.superscript = True
authors.add_run(', ' + names[1] + ',\n' + names[2] + ' and ' + names[3])
address = paragraph('', 'address')
address.add_run('* Corresponding author: ')
address.add_run(emails[0]).style = doc.styles['e-mail']
address.add_run('\n')
address.add_run(emails[1] + '; ' + emails[2]).style = doc.styles['e-mail']
address.add_run('\n')
address.add_run(emails[3]).style = doc.styles['e-mail']
abstract = ' '.join(text(original[i]) for i in [4, 5, 6]).removeprefix('Abstract—')
assert len(abstract.split()) <= 250
paragraph('Abstract. ' + abstract, 'abstract', bold_prefix='Abstract. ')
paragraph('Keywords: ' + text(original[7]).removeprefix('Keywords—'), 'keywords', bold_prefix='Keywords: ')

section_names = ['Introduction', 'Literature Review', 'Comparative Study',
                 'Proposed Methodology', 'Results and Discussion', 'Future Scope', 'Conclusion']
major_indices = {8: 0, 19: 1, 34: 2, 39: 3, 77: 4, 105: 5, 108: 6}
subheading_only = {40, 42, 70, 78, 86, 92, 100}
inline_subheadings = {20, 25, 29, 30, 48, 59}
first_paragraph = True
for index in range(8, 111):
    element = original[index]
    value = text(element)
    roman = {'I':'1', 'II':'2', 'III':'3', 'IV':'4', 'V':'5', 'VI':'6', 'VII':'7'}
    value = re.sub(r'\bSection (VII|VI|IV|III|II|V|I)\b', lambda m:'Section '+roman[m.group(1)], value)
    if index in major_indices:
        heading(section_names[major_indices[index]])
        first_paragraph = True
    elif index in subheading_only:
        heading(re.sub(r'^[A-Z]\.\s*', '', value), 2)
        first_paragraph = True
    elif index in inline_subheadings:
        label, remainder = value.split(':', 1)
        heading(re.sub(r'^[A-Z]\.\s*', '', label), 2)
        paragraph(remainder.strip(), 'p1a')
        first_paragraph = False
    elif element.tag == qn('w:tbl'):
        copy_table(element); first_paragraph = True
    elif element.xpath('.//m:oMathPara'):
        copy_math(element); first_paragraph = True
    elif element.xpath('.//w:drawing'):
        copy_figure(element); first_paragraph = True
    elif value.startswith('TABLE '):
        number = {32: 1, 57: 2, 90: 3}[index]
        value = re.sub(r'^TABLE [IVX]+\.', f'Table {number}.', value)
        p = paragraph(value, 'tablecaption', bold_prefix=f'Table {number}.')
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif value.startswith('Fig. '):
        prefix = re.match(r'Fig\. \d+\.', value).group()
        p = paragraph(value, 'figurecaption', bold_prefix=prefix)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        first_paragraph = True
    elif value:
        if index == 18:
            value = value.replace('Section II', 'Section 2').replace('Section III', 'Section 3').replace('Section IV', 'Section 4').replace('Section V', 'Section 5')
        value = value.replace('Table III', 'Table 3')
        paragraph(value, 'p1a' if first_paragraph else 'Normal')
        first_paragraph = False

# Keep the introduction of each displayed equation on the same page as the math.
flow = list(doc._element.body)
for current, following in zip(flow, flow[1:]):
    if current.tag == qn('w:p') and following.xpath('.//m:oMathPara'):
        Paragraph(current, doc._body).paragraph_format.keep_with_next = True

paragraph('Disclosure of Interests. The authors have no competing interests to declare.',
          'acknowlegments', bold_prefix='Disclosure of Interests. ')
p = paragraph('References', 'heading1'); no_number(p)

# The bibliography's metadata is retained; only the publisher's presentation changes.
references = [
'Dong, H., Blunt, M.J.: Pore-network extraction from micro-computerized-tomography images. Physical Review E 80, 036307 (2009). https://doi.org/10.1103/PhysRevE.80.036307',
'Gostick, J., et al.: OpenPNM: A pore network modeling package. Computing in Science & Engineering 18(4), 60–74 (2016). https://doi.org/10.1109/MCSE.2016.49',
'Wu, J., Yin, X., Xiao, H.: Seeing permeability from images: Fast prediction with convolutional neural networks. Science Bulletin 63(18), 1215–1222 (2018). https://doi.org/10.1016/j.scib.2018.08.006',
'Hamilton, W.L., Ying, R., Leskovec, J.: Inductive representation learning on large graphs. Advances in Neural Information Processing Systems 30 (2017). arXiv:1706.02216',
'Chapuis, R.P., Aubertin, M.: On the use of the Kozeny–Carman equation to predict the hydraulic conductivity of soils. Canadian Geotechnical Journal 40(3), 616–628 (2003). https://doi.org/10.1139/t03-013',
'Qian, Y.H., d’Humières, D., Lallemand, P.: Lattice BGK models for Navier–Stokes equation. Europhysics Letters 17(6), 479–484 (1992). https://doi.org/10.1209/0295-5075/17/6/001',
'Savalkar, V.A., Saini, G.S., Pawar, S.D.: Density-Aware Explainable Hybrid CNN–Transformer Framework for Automated BI-RADS Breast Density Classification in Mammographic Images. Journal of Intelligent Decision Making and Information Science 3(1s), 1003–1027 (2026). https://doi.org/10.59543/jidmis.v3.480',
'Alzahrani, M.K., Shapoval, A., Chen, Z., et al.: Pore-GNN: A graph neural network-based framework for predicting flow properties of porous media from micro-CT images. Advances in Geo-Energy Research 10(1), 39–55 (2023). https://doi.org/10.46690/ager.2023.10.05',
'Zhao, X., Zhong, Y., Li, P.: RTG-GNN: A novel rock topology-guided approach for permeability prediction using graph neural networks. Geoenergy Science and Engineering 243, 213358 (2024). https://doi.org/10.1016/j.geoen.2024.213358',
'I’Anson, J.M., Simmons, M.J.H., Stitt, E.H., Gallen, R.W.: Using graph neural networks to predict the permeability of porous media. InterPore Journal 2(3) (2025). https://doi.org/10.69631/ipj.v2i3nr47',
'Zhao, Q., Xiao, H.: An end-to-end differentiable, graph neural network-embedded pore network model for permeability prediction. Journal of Hydrology 675, 135602 (2026). https://doi.org/10.1016/j.jhydrol.2026.135602',
'Imperial College London, Pore-scale Modelling Group: ImagesICPSC2007. Binary images and source headers. https://imperialcollegelondon.box.com/v/ImagesICPSC2007, last accessed 2026/10/22',
'Imperial College London: Estaillades 1000³ micro-CT image, 3.31136 μm. Official archive: https://imperialcollegelondon.box.com/shared/static/2n4lb95amz7fstk19q826uzu922eb4b1.zip, last accessed 2026/10/22',
'Gostick, J.T.: Versatile and efficient pore network extraction method using marker-based watershed segmentation. Physical Review E 96, 023307 (2017). https://doi.org/10.1103/PhysRevE.96.023307',
'Veličković, P., et al.: Graph Attention Networks. International Conference on Learning Representations (2018). arXiv:1710.10903',
'Sherman, J., Morrison, W.J.: Adjustment of an inverse matrix corresponding to a change in one element of a given matrix. The Annals of Mathematical Statistics 21(1), 124–127 (1950). https://doi.org/10.1214/aoms/1177729893',
'OpenPNM: Valvatne–Blunt hydraulic conductance. Official implementation documentation, installed version 3.6.0. https://openpnm.org/_modules/openpnm/models/physics/hydraulic_conductance/_funcs.html#valvatne_blunt, last accessed 2026/10/22',
'Nassiri, S., van Paassen, L., Mahabadi, N.: Evolution of porous media and porosity–permeability relations during biogeochemical clogging processes: A pore-scale simulation study. Computers and Geotechnics 198, 108292 (2026). https://doi.org/10.1016/j.compgeo.2026.108292'
]
for reference in references:
    p = paragraph(reference, 'referenceitem')
    p.paragraph_format.keep_together = True

# Replace only the running sample text, leaving the native PAGE fields intact.
for part in doc.part.package.parts:
    if str(part.partname).startswith('/word/header'):
        for node in part.element.iter(qn('w:t')):
            if node.text == 'F. Author and S. Author': node.text = 'S. Rane et al.'
            if node.text == 'Contribution Title (shortened if too long)':
                node.text = 'Graph Attention for Permeability and Throat Discovery'
for footnote in doc.part._element.xpath('.//w:footnoteReference'):
    raise AssertionError('Unexpected reference to a template footnote')
settings = doc.settings.element
for tag in ['documentProtection', 'writeProtection']:
    for element in settings.findall(qn('w:' + tag)):
        settings.remove(element)
update = settings.find(qn('w:updateFields'))
if update is None:
    update = OxmlElement('w:updateFields'); settings.append(update)
update.set(qn('w:val'), 'true')
# Sample chart and disclosure footnote are not part of the manuscript.
for rid, rel in list(doc.part.rels.items()):
    if rel.reltype.endswith('/chart'):
        doc.part.drop_rel(rid)
for part in doc.part.package.parts:
    if str(part.partname) == '/word/footnotes.xml':
        notes = E.fromstring(part.blob)
        for note in list(notes):
            if int(note.get(qn('w:id'), '0')) > 0:
                notes.remove(note)
        part._blob = E.tostring(notes, xml_declaration=True, encoding='UTF-8', standalone=True)

doc.core_properties.title = title
doc.core_properties.author = '; '.join(names)
doc.core_properties.subject = 'TopoFlow conference manuscript in the Atlantis Press linked proceedings template'
doc.core_properties.comments = 'Formatting conversion of completed experiments; original manuscript retained.'
assert E.tostring(doc._element.body.sectPr) == section_xml
assert len(doc.inline_shapes) == 4 and len(doc.tables) == 3
assert len(doc._element.xpath('//m:oMathPara')) == 11
doc.save(OUTPUT)
# python-docx can reserialize even untouched XML parts. Restore the template's
# preserve-only parts verbatim after the object-model edit.
preserve_parts = ['word/theme/theme1.xml', 'word/fontTable.xml', 'word/numbering.xml',
                  'word/webSettings.xml', 'word/endnotes.xml', 'word/footer1.xml',
                  'word/footer2.xml', 'word/footer3.xml', 'customXml/item1.xml',
                  'customXml/itemProps1.xml', 'customXml/_rels/item1.xml.rels']
buffer = BytesIO()
with ZipFile(OUTPUT) as result_zip, ZipFile(REFERENCE) as template_zip, ZipFile(buffer, 'w', ZIP_DEFLATED) as target_zip:
    for item in result_zip.infolist():
        data = template_zip.read(item.filename) if item.filename in preserve_parts else result_zip.read(item.filename)
        target_zip.writestr(item, data)
OUTPUT.write_bytes(buffer.getvalue())
assert sha(SOURCE) == 'd9b7f24f5f74370cf1491e4172ec5a9e750cba27643fb9a6c4b1002c4ce281cf'
with ZipFile(REFERENCE) as z:
    inventory = {n: {'bytes':len(z.read(n)), 'sha256':hashlib.sha256(z.read(n)).hexdigest()} for n in z.namelist()}
(WORK / 'template/package_inventory.json').write_text(json.dumps(inventory, indent=2), encoding='utf-8')
proof = {'source':SOURCE.name,'source_sha256':sha(SOURCE), 'template_sha256':sha(REFERENCE),
         'output':OUTPUT.name,'output_sha256':sha(OUTPUT),'abstract_words':len(abstract.split()),
         'native_equations':11,'figures':4,'data_tables':3,'references':18,'sections':1,
         'corresponding_author':'Shaunak Rane','competing_interests_confirmed_by_user':True}
(WORK / 'conversion_audit.json').write_text(json.dumps(proof, indent=2), encoding='utf-8')
print(json.dumps(proof, indent=2))
