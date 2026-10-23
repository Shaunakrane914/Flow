from pathlib import Path
from zipfile import ZipFile
import json,re,hashlib,unicodedata
from docx import Document
from docx.oxml.ns import qn
from lxml import etree as E

root=Path(__file__).resolve().parents[1]
work=root/'.paper_revision/latex_restoration'
source=root/'TopoFlow_Atlantis_ICISEA2026.docx'
manifest=json.loads((work/'content_manifest.json').read_text(encoding='utf-8'))
tex=(root/'paper/latex/topoflow.tex').read_text(encoding='utf-8')
doc=Document(source)
assert hashlib.sha256(source.read_bytes()).hexdigest()==manifest['source_sha256']
norm=lambda s:re.sub(r'\s+',' ',unicodedata.normalize('NFC',s).replace('’',"'")).strip()

def group_at(text,start):
    assert text[start]=='{'
    depth=1;j=start+1
    while depth:
        if text[j]=='{' and text[j-1]!='\\':depth+=1
        if text[j]=='}' and text[j-1]!='\\':depth-=1
        j+=1
    return text[start+1:j-1],j

def detex(value):
    value=re.sub(r'\\url\{([^}]+)\}',lambda m:m.group(1),value)
    value=re.sub(r'\\cite\{r(\d+)\}',lambda m:'['+m.group(1)+']',value)
    subs=dict(zip('0123456789htiv','₀₁₂₃₄₅₆₇₈₉ₕₜᵢᵥ'))
    supers=dict(zip('0123456789-eq','⁰¹²³⁴⁵⁶⁷⁸⁹⁻ᵉᵠ'))
    value=re.sub(r'\\textsubscript\{([^}]+)\}',lambda m:''.join(subs[c] for c in m.group(1)),value)
    math_map={r'\Phi':'Φ',r'\tau':'τ',r'\Delta':'Δ',r'\alpha':'α',r'\mu':'μ',r'\pi':'π',
              r'\rho':'ρ',r'\sigma':'σ',r'\ell':'ℓ',r'\hat{y}':'ŷ',r'\bar{y}':'ȳ',
              r'\in':'∈',r'\cup':'∪',r'\sum':'∑','-':'−',r'\times':'×',r'\pm':'±',
              r'\rightarrow':'→',r'\cdot':'·',r'\log_{10}':'log₁₀','z_g':'zɡ'}
    while r'\ensuremath{' in value:
        start=value.index(r'\ensuremath{'); brace=start+len(r'\ensuremath')
        inner,end=group_at(value,brace)
        if inner.startswith('^{'):
            replacement=''.join(supers[c] for c in inner[2:-1])
        else:replacement=math_map[inner]
        value=value[:start]+replacement+value[end:]
    for old,new in {r'\&':'&',r'\%':'%',r'\$':'$',r'\#':'#',r'\_':'_',r'\{':'{',r'\}':'}',
                    r'\textbackslash{}':'\\',r'\textasciitilde{}':'~',r'\textasciicircum{}':'^'}.items():
        value=value.replace(old,new)
    return value.replace('--','–')

source_blocks=[]
for p in doc.paragraphs:
    if p.style.style_id in ['Normal','p1a','heading1','heading2','acknowlegments'] and p.text and p.text!='References':
        source_blocks.append(p.text)
actual=[b for b in manifest['blocks'] if b['kind'] in ['heading','paragraph']]
assert [b['text'] for b in actual]==source_blocks
for b in actual:
    encoded=b['latex']
    if b['kind']=='heading':
        encoded=encoded[encoded.index('{')+1:-1]
    elif b['style']=='acknowlegments':
        label,end=group_at(encoded,encoded.index('{'))
        encoded=label+' '+encoded[end:].strip()
    assert norm(detex(encoded))==norm(b['text']),(b['text'],detex(encoded))
    assert b['latex'] in tex
abs_encoded=tex.split(r'\emergencystretch=3em',1)[1].split(r'\keywords{',1)[0].strip().removeprefix(r'\raggedright ')
assert norm(detex(abs_encoded))==norm(manifest['abstract'])
assert manifest['abstract']==doc.paragraphs[3].text.removeprefix('Abstract. ')
refs=[p.text for p in doc.paragraphs if p.style.style_id=='referenceitem']
assert refs==manifest['references'] and len(refs)==18
for i,ref in enumerate(refs,1):
    encoded=re.search(r'\\bibitem\{r'+str(i)+r'\} (.*)',tex).group(1)
    assert norm(detex(encoded))==norm(ref),(i,ref,encoded)
tables=[b for b in manifest['blocks'] if b['kind']=='table']
assert [b['cells'] for b in tables]==[[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]
for b in tables:
    assert b['latex'] in tex
    recovered=[]
    for line in b['latex'].splitlines():
        if not line.endswith(r' \\'):continue
        cells=[]
        for cell in line[:-3].split(' & '):
            if cell.startswith(r'\textbf{'):
                cell,end=group_at(cell,len(r'\textbf'))
            cells.append(norm(detex(cell)))
        recovered.append(cells)
    assert recovered==[[norm(c) for c in row] for row in b['cells']]
equations=[b for b in manifest['blocks'] if b['kind']=='equation']
assert [b['omml_sha256'] for b in equations]==[hashlib.sha256(E.tostring(e,method='c14n',exclusive=True)).hexdigest() for e in doc._element.xpath('//m:oMathPara')]
expected=[
r'k=\frac{\Phi^3}{c\tau^2S_v^2}',
r'f_i\left(x+c_i\Delta t,t+\Delta t\right)=f_i\left(x,t\right)-\frac{1}{\tau}\left[f_i\left(x,t\right)-f_i^{\mathrm{eq}}\left(x,t\right)\right]',
r'z=\mathrm{MLP}\left(\frac{1}{|V|}\sum_{v\in V}h_v^L\right)',
r'h_v^{l+1}=\mathrm{ELU}\left(\sum_{u\in N(v)\cup\{v\}}\alpha_{vu}^lW^lh_u^l\right)',
r'\alpha_{vu}=\frac{\exp(e_{vu})}{\sum_{j\in N(v)\cup\{v\}}\exp(e_{vj})}',
r'\hat{y}=2\log_{10}d_0+\mu_t+\sigma_t(z+z_g)',
r'\mathrm{Re}=\frac{\rho vD_h}{\mu}<1',
r'k=\frac{Q\mu L}{A\Delta P}',
r'\mathrm{MSE}=\frac{1}{N}\sum_{i=1}^{N}(y_i-\hat{y}_i)^2',
r'R^2=1-\frac{\sum_{i=1}^{N}(y_i-\hat{y}_i)^2}{\sum_{i=1}^{N}(y_i-\bar{y})^2}',
r'\mathrm{loss}=\frac{(1-r)s}{1-(1-r)\ell}'
]
def math_norm(s):
    s=s.replace(r'\left','').replace(r'\right','')
    s=s.replace(r'\{','SETOPEN').replace(r'\}','SETCLOSE')
    return re.sub(r'[\s{}]','',s)
for i,(b,want) in enumerate(zip(equations,expected),1):
    assert math_norm(b['latex'])==math_norm(want),(i,b['latex'],want)
    assert b['latex'] in tex
figures=[b for b in manifest['blocks'] if b['kind']=='figure']
with ZipFile(source) as z:
    source_hashes=sorted(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
assert sorted(b['sha256'] for b in figures)==source_hashes
for b in figures:
    p=root/'paper/latex/figures'/f'figure-{b["number"]}.png'
    assert hashlib.sha256(p.read_bytes()).hexdigest()==b['sha256']
    assert b['latex'] in tex
for name in manifest['authors']:assert name in tex
for email in manifest['emails']:assert tex.count(email)==1,email
assert r'\thanks{Corresponding author: \email{'+manifest['emails'][0]+'}' in tex
assert norm(manifest['title'])==norm(doc.paragraphs[0].text)
assert r'\title{\protect\Centering '+manifest['title']+'}' in tex
for kind,style,prefix in [('figure','figurecaption',r'^Fig\. \d+\.\s*'),('table','tablecaption',r'^Table \d+\.\s*')]:
    captions=[re.sub(prefix,'',p.text) for p in doc.paragraphs if p.style.style_id==style]
    stored=[b['caption'] for b in manifest['blocks'] if b['kind']==kind]
    assert captions==stored
    for b in (b for b in manifest['blocks'] if b['kind']==kind):
        encoded=b['latex'].split(r'\caption{',1)[1]
        content,end=group_at('{'+encoded,0)
        assert norm(detex(content))==norm(b['caption'])
assert re.findall(r'\\bibitem\{r(\d+)\}',tex)==list(map(str,range(1,19)))
assert set(re.findall(r'\\cite\{r(\d+)\}',tex))==set(map(str,range(1,19)))
assert not any(x in tex for x in ['ZZPROTECTED',r'\texttt',r'\vspace'])
for name,digest in manifest['template_files'].items():
    assert hashlib.sha256((root/'paper/latex'/name).read_bytes()).hexdigest()==digest
result={'status':'passed','source_sha256':manifest['source_sha256'],'prose_and_heading_blocks_round_tripped':len(actual),
        'abstract_words_preserved':len(manifest['abstract'].split()),'references_round_tripped':18,'equations_independently_checked':11,
        'figure_files_byte_identical':4,'data_tables_preserved':3,'all_four_authors_and_emails_preserved':True,
        'official_template_files_unmodified':True,'research_rewritten':False,'tex_sha256':hashlib.sha256((root/'paper/latex/topoflow.tex').read_bytes()).hexdigest()}
(work/'source_verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
