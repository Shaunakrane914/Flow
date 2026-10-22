"""Targeted revision of the retained IEEE Word paper using committed evidence.

Run with the Codex managed Python runtime. No experiments or application code
are changed. The source DOCX is retained; paragraphs are edited in place.
"""
from copy import deepcopy
import csv
import hashlib
import json
import math as pymath
import re
import subprocess
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from lxml import etree as E
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / '.paper_revision/current'
WORK.mkdir(parents=True, exist_ok=True)
SOURCE = ROOT / 'TopoFlow_IEEE_V4_Cited_Equations.docx'
OUTPUT = ROOT / 'TopoFlow_IEEE_V4_Cited_Revised.docx'
doc = Document(SOURCE)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == '0297f6ad073d2c5003d713b63f36a27d2bbb36e912f5fa2ba950f08485b4da9c'
body = list(doc._element.body)
assert len(body) == 165 and len(doc.inline_shapes) == 4, 'Source document changed'
assert 'Millisecond' in Paragraph(body[0], doc._body).text
original_author = E.tostring(body[1])
sections_before = [E.tostring(e) for e in doc._element.xpath('//w:sectPr')]
changes = []

def para(index):
    assert body[index].tag == qn('w:p'), index
    return Paragraph(body[index], doc._body)

def replace(index, text, bold_prefix=None):
    p = para(index)
    before = p.text
    props = next((deepcopy(r._r.rPr) for r in p.runs if r._r.rPr is not None), None)
    p.clear()
    if bold_prefix and text.startswith(bold_prefix):
        chunks = [(bold_prefix, True), (text[len(bold_prefix):], False)]
    else:
        chunks = [(text, None)]
    for chunk, bold in chunks:
        r = p.add_run(chunk)
        if props is not None:
            r._r.insert(0, deepcopy(props))
        if r._r.rPr is not None:
            for prop in ['vertAlign','position','vanish','smallCaps','caps']:
                for node in r._r.rPr.findall(qn('w:'+prop)):
                    r._r.rPr.remove(node)
        r.bold = bold if bold is not None else False
    changes.append({'source_body_index': index, 'before': before, 'after': text})
    return p

def after(index, text, heading=False):
    e = OxmlElement('w:p')
    body[index].addnext(e)
    p = Paragraph(e, doc._body)
    p.style = doc.styles['Heading 2'] if heading else doc.styles['Normal']
    p.add_run(text)
    if heading:
        p.paragraph_format.keep_with_next = True
    return p

def heading(index, text):
    p = replace(index, text)
    p.style = doc.styles['Heading 1']
    pr = p._p.get_or_add_pPr()
    for n in pr.findall(qn('w:numPr')):
        pr.remove(n)
    num = OxmlElement('w:numPr')
    ident = OxmlElement('w:numId'); ident.set(qn('w:val'), '0')
    num.append(ident); pr.append(num)
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    return p

# Original section order and subject are retained. Numbers now follow that order.
title = replace(0, 'Permeability Prediction in Micro CT Scans using Graph Attention Networks')
if 'Title' not in doc.styles:
    doc.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)
title.style = doc.styles['Title']
for r in title.runs:
    r.font.name = 'Times New Roman'; r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0, 0, 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, text in [(8,'I. INTRODUCTION'), (19,'II. LITERATURE REVIEW'),
                (45,'III. COMPARATIVE STUDY'), (51,'IV. PROPOSED METHODOLOGY'),
                (93,'V. RESULTS AND DISCUSSION'), (111,'VI. FUTURE SCOPE'),
                (120,'VII. CONCLUSION'), (124,'REFERENCES')]:
    heading(i, text)

edits = {
4: 'Abstract—Pore-network graphs offer a compact representation for predicting permeability from micro-CT images, but accuracy on simulation labels does not establish transfer to unseen rock structures. We evaluate TopoFlow, a three-layer graph attention network (GAT), using the Imperial College ImagesICPSC2007 collection and the Estaillades carbonate scan at 3.31136 μm voxel resolution.',
5: 'The workflow retains SNOW2 extraction, graph message passing and permeability regression. Four development scans provide 24 natural-rock crops for nested validation; a locked test uses 32 crops from four further scans. Targets are cylindrical-throat pore-network-model (PNM) calculations. Scale normalization and graph descriptors improve GAT performance, but the primary GAT attains external log-RMSE 0.569 compared with 0.253 for a physics ridge baseline.',
6: 'A separate frozen evaluation on ten crops from five previously unused scans tests geometry-only critical-throat selection. Individually restricting GAT-selected throats by 50% yields mean permeability loss of 3.57%, versus 0.37% for random selection and 4.08% for an affine-pressure ranking. GAT response MAE is 2.04 percentage points, versus 0.46 for no response. GAT identified more consequential throats than random selection, but affine-pressure ranking performed better overall and response magnitudes remained poorly calibrated.',
7: 'Keywords—Graph attention networks, pore-network modeling, permeability, digital rock physics, SNOW2, critical throats.',
9: 'Accurate prediction of absolute permeability in heterogeneous porous structures is important in groundwater transport, hydrocarbon recovery and geological carbon storage. Micro-CT imaging resolves pore-scale geometry, while the relationship between that geometry and macroscopic flow depends on connectivity, throat dimensions and boundary conditions.',
10: 'Direct numerical methods such as the lattice Boltzmann method (LBM) resolve flow on voxel grids. Pore-network models instead approximate the connected void space by pores and throats [2], [4]. The resulting permeability depends on both the extraction method and the hydraulic closure. A learned surrogate trained on these labels therefore approximates a specified network calculation.',
11: 'Data-driven models, including three-dimensional convolutional neural networks (CNNs), have been used to predict permeability from images [6]. Their evaluation requires source-separated tests because nearby crops share texture and pore structure. A random crop split can give an optimistic account of generalization to another scan.',
12: 'Graph representations process extracted pores and their connections, avoiding feature propagation through every solid voxel. Network extraction still incurs a measurable cost. We compare prediction accuracy and computational cost with analytical relations, descriptor models and the reference network solver.',
13: 'TopoFlow retains this image-to-network approach and uses graph attention for permeability regression. We investigate how physical scale normalization and graph descriptors affect transfer across source scans. A follow-up examines localized throat restrictions, which a model receiving only pore diameter, pore volume and connectivity cannot observe when those inputs remain unchanged.',
14: 'The study makes three empirical contributions:',
15: '1) A source-separated comparison of GAT, GraphSAGE [15] and simpler permeability predictors, with nested model selection, five training seeds and a locked external test.',
16: '2) A controlled ablation separating scale normalization from graph descriptors, followed by a frozen, geometry-only critical-throat evaluation with matched learned and analytical controls.',
17: '3) Component-level CPU timing and hydraulic-reference checks that distinguish learning error, numerical consistency and uncertainty in the pore-network closure.',
18: 'Section II reviews related methods; Section III defines the comparison. Section IV presents the workflow and mathematical model. Section V reports completed experiments, followed by future scope and conclusions.',
20: 'A. Physics Based Analytical Models: Traditional permeability estimates include the Kozeny–Carman (KC) relation [1]. This family relates porosity, specific surface and tortuosity to transport through an idealized medium. Its low computational cost makes it a necessary baseline even when the pore structure is heterogeneous.',
21: 'Chapuis and Aubertin discuss the applicability of KC to soil hydraulic conductivity [1]. Such relations do not explicitly resolve individual flow paths, and their performance depends on the chosen scale and calibration. We evaluate a calibrated porosity relation and a diameter-scaled variant.',
22: 'A conventional form of the KC relation is',
24: 'where Φ is porosity, τ is tortuosity, c is a dimensionless geometric constant and Sᵥ is specific surface with inverse-length units. The calibrated benchmark variants are defined separately in Section III.',
25: 'B. Lattice Boltzmann Simulation: LBM advances discrete particle distributions on a lattice and can resolve voxel-scale flow. No LBM labels are generated in the present experiments. For background, we use the lattice-BGK formulation of Qian et al. [16].',
26: 'Its single-relaxation-time update is',
28: 'Here fᵢ is the particle distribution for lattice direction i, cᵢ is its discrete velocity, Δt is the time step, fᵢᵉᵠ is the equilibrium distribution and τ is the relaxation time. This background equation is distinct from the hydraulic PNM used for our targets.',
29: 'C. Image Based Machine Learning: Wu, Yin and Xiao predict permeability with CNNs [6]. Image models establish that geometry can support regression, but comparisons across architectures require matching data, labels and timing boundaries. We do not report a CNN runtime or memory comparison because no matched CNN experiment was performed.',
30: 'D. Graph Neural Networks: Pore-GNN combines CNN-derived image features with graph models for permeability prediction [7]. These feature graphs differ from physical pore–throat networks. RTG-GNN incorporates pore and throat information with gating and attention [8]. I’Anson et al. evaluate graph permeability prediction against LBM labels and include extraction cost in their assessment [9]. Graph-based permeability prediction therefore has several established representations and training approaches.',
31: 'Zhao and Xiao embed learned conductances in a differentiable PNM [10]. Building on established graph-based permeability prediction, our comparison focuses on source-scan transfer, the separate effects of scale and descriptors, and geometry-only selection of influential throats under a fixed hydraulic closure.',
43: 'TABLE I. Scope of permeability estimation approaches',
46: 'The permeability benchmark compares a common log-permeability target under identical source grouping. The node-only GAT is compared with GraphSAGE, uniform attention, a pore-set model without message passing, descriptor-based ExtraTrees, a ridge model and two calibrated KC variants. Hyperparameters and epoch counts are selected using development parents only.',
47: 'The porosity KC baseline fits an intercept to log₁₀[Φ³/(1−Φ)²]. The scaled variant additionally includes twice the log geometric-mean pore diameter. Physics ridge regresses the scale-adjusted target on the graph descriptors with regularization selected on development scans. These are fitted predictors rather than parameter-free evaluations of the classical KC equation.',
48: 'For critical-throat selection, the guided GAT is compared with a guided mean-aggregation GNN, a guided edge MLP and an unguided GAT. All use the same geometry inputs and bounded response readout; the guided controls share the importance supervision. The mean model is GraphSAGE-style, whereas the bulk benchmark uses the GraphSAGE implementation directly.',
49: 'The analytical throat rankings use conductance proportional to d⁴/L and an affine-pressure proxy proportional to g(Δx)² along the flow direction. They require geometry only. A PNM-sensitivity ranking and exact finite-change updates use solved flow quantities and are reported as solver-based references. No-response prediction is included as a response-accuracy control.',
54: 'TopoFlow converts segmented micro-CT images to pore-network graphs and predicts absolute permeability. The main pipeline remains SNOW2 extraction, three GAT layers, global mean pooling and an MLP readout. The experiments distinguish bulk regression from an edge-aware extension for single-throat restrictions; their datasets and evaluation endpoints are reported separately.',
52: 'A. Overview',
56: 'B. System Architecture',
57: '1) Image input: Natural-rock binary volumes come from the Imperial College collection [11] and the separate Estaillades archive [12]. Fixed 128³-voxel crops retain the physical voxel size from source metadata. Pore space is encoded by zero. Source-byte checks verify crop provenance and exclude duplicate test crops.',
58: '2) Network extraction: SNOW2 partitions the void space using watershed regions [3]. Pore bodies form vertices and connecting throats form edges. The fixed settings are sigma 0.4, maximum radius 4, boundary width (3,0,0) voxels and one extraction division. Geometry is converted to physical units before the hydraulic reference is calculated.',
59: '3) Graph prediction: The bulk encoder receives log pore diameter and log pore volume. Optional physical normalization and ten graph descriptors are added in controlled variants. The model predicts log₁₀ permeability in m²; conversion to linear units is performed only after reconstructing the target scale.',
62: 'Fig. 1. TopoFlow workflow. The PNM supplies training and evaluation targets separately from model inference.',
63: 'C. Model Description: The bulk model uses graph attention [5]. Its first layer has two 64-channel heads, giving 128 channels; the next two have one 64-channel head each. ELU and dropout follow each layer. Global mean pooling feeds a 64→32→1 MLP with ReLU and dropout. The original bulk encoder does not receive throat features.',
64: 'The pooled GAT output is a scalar z used for standardized-target regression:',
66: 'Each attention layer updates a pore embedding by a weighted neighborhood sum:',
68: 'The coefficients α are softmax-normalized over the neighbors and self-loop of pore v. Scores use a learned vector applied to concatenated transformed endpoint embeddings followed by LeakyReLU [5]. The two first-layer heads are concatenated. This local attention is distinct from the separately trained global throat-importance head used in the restriction experiment.',
69: 'For the scale-normalized variant, d₀ is the geometric-mean diameter of physical pores. Inputs become log₁₀(d/d₀) and log₁₀(V/d₀³); twice log₁₀d₀ is subtracted from the target. Training-only means and standard deviations normalize features and targets. The geometry branch uses a 10→16→1 MLP with scalar output zɡ. Prediction reconstructs the log target using the development residual mean μₜ and standard deviation σₜ:',
72: 'TABLE II. Geometry used by the bulk and restriction models',
73: 'D. Mathematical Details: The hydraulic reference assumes steady, incompressible creeping flow. Cylindrical throat conductance is g = πd⁴/(128μL), using inscribed diameter d, direct pore-center distance L and viscosity μ = 0.001 Pa·s. Inlet and outlet pressures are 1 and 0 Pa. Nonspanning components are excluded and conservation residuals checked.',
74: 'A conventional creeping-flow criterion is',
76: 'where ρ is density, v is characteristic velocity and Dₕ is hydraulic diameter. This criterion states the assumed flow regime; it is not a measured Reynolds number for every image. Effective permeability is obtained from the steady network flux through Darcy’s relation:',
78: 'Here Q is total volumetric flow rate, L is the physical image length in the flow direction, A is the original image cross-sectional area and ΔP is the applied pressure difference. Artificial boundary conduits are included in the declared reference. The labels are PNM approximations, not LBM or measured ground truth [4].',
79: 'Training uses mean squared error in y = log₁₀(k/m²), or its training-standardized residual in the scale-normalized model:',
81: 'Predictive R² is calculated in log space using the mean reference value of the evaluated set:',
83: 'Here N is the number of evaluated crops, yᵢ is reference log permeability and ŷᵢ is its prediction. RMSE, MAE and R² are recomputed from saved predictions. Negative test R² is retained. Parent-balanced RMSE and pooled crop metrics are distinguished; seed variability is not treated as independent geological replication.',
84: 'E. Working Steps',
85: '1) Input and separation: The development set contains 24 natural crops from Berea (8), C1 (7), C2 (6) and Estaillades (3). Outer validation holds out one complete parent scan, with inner selection on the other three. No crop from an outer test parent is used for feature scaling, target scaling or model selection.',
86: '2) Locked external test: After development selection, models are trained on all 24 development crops and tested on 32 crops, eight each from S1, S3, S5 and S9. Their voxel sizes are 8.683, 9.100, 3.997 and 3.398 μm. The primary GAT variant and its settings are fixed before inspecting these outcomes.',
88: '3) Feature assignment: The ten bulk graph descriptors comprise two porosity transforms, pore log-diameter spread, normalized log-volume mean and spread, log pore and throat counts, and mean, standard deviation and coefficient of variation of degree. Ablations independently enable normalization and this descriptor branch.',
90: '4) Training: The neural benchmark uses seeds 11, 22, 33, 44 and 55. The locked primary GAT uses learning rate 0.003, dropout 0.1, weight decay 0.001 and 14 epochs selected on development scans. GAT and GraphSAGE receive matching data and selection procedures. The graph-wise loss pairs one scalar target with one prediction per graph.',
92: '5) Restriction extension: A separate edge-aware model trains on eight development crops, two per parent. Seven node features add normalized coordinates and inlet/outlet flags; six edge features encode normalized diameter, length, a d⁴/L proxy and absolute direction components. Three 32-channel GAT layers feed a symmetric 70→32→2 edge head. A global softmax predicts importance s, and a sigmoid gives leverage ℓ bounded by 1−s. For diameter reduction q, the conductance multiplier is r=(1−q)⁴. Changing one conductance gives a rank-one update of the linear network equations. The Sherman–Morrison identity [17] yields the loss expression below when s and ℓ are the exact dissipation fraction and edge leverage. The GAT substitutes learned estimates; their accuracy determines response calibration. A derivation accompanies the source implementation:',
94: 'A. Permeability Prediction',
95: 'The evaluation uses natural-rock crops with parent-scan separation. Synthetic development examples and historical demonstration outputs are excluded from the reported benchmark. The hydraulic target is the same cylindrical-throat PNM for all compared bulk predictors.',
96: 'Nested validation gives parent-balanced log-RMSE 0.8311 ± 0.0200 for the normalized GAT with descriptors and 0.8713 ± 0.0781 for matched GraphSAGE. Calibrated porosity KC gives 0.7525. The GAT improvement over GraphSAGE is modest and does not establish architectural superiority on four source groups.',
97: 'The controlled GAT ablation yields parent-balanced log-RMSE 1.5528 for standard inputs, 0.9504 with normalization alone, 0.8933 with descriptors alone and 0.8311 with both. Thus, both additions help in development validation. This comparison separates feature and scale choices from an attention-specific explanation.',
98: 'On the locked external test, the primary normalized GAT with descriptors gives log-RMSE 0.5687 ± 0.0805 and log-R² −0.7671. Physics ridge gives 0.2534 and 0.6547, respectively. The two single-addition GAT ablations obtain RMSE 0.4111 and 0.4131, but they are not retrospectively substituted for the primary model. Table III includes the fixed comparison.',
99: 'Figure 2 summarizes external errors. Values for neural models are means over five independently trained seeds; the error bars show seed standard deviation. The simpler ridge and scaled-KC controls demonstrate that low-dimensional geometric information explains substantial variation in the reference labels. An accurate node-only predictor is therefore not, by itself, evidence that attention has learned the governing flow paths.',
101: 'B. Computational Cost: Serial CPU measurements use one Torch/BLAS thread. Bulk-model forward timing uses ten warmups and thirty calls per crop, excluding file loading and feature normalization. Median GAT forward time is 1.72 ms across 24 crops. Separate measurements on one crop per development parent give median SNOW2 extraction 4.06 s and PNM solve 1.85 ms. These samples and operation scopes differ; their ratio is not an end-to-end speedup.',
103: 'Fig. 3. Measured component costs in the development benchmark. Extraction is measured on four crops and neural forwarding on 24; file loading and model loading are excluded.',
104: 'TABLE III. Locked external permeability results on 32 crops',
105: 'C. Critical Throat Evaluation',
106: 'Partial throat narrowing leaves the original node-only GAT inputs unchanged, so that model returns exactly the same permeability. The edge-aware extension addresses this representation limitation. Guided GAT, guided mean GNN and guided edge MLP share importance supervision with weight 0.05; an unguided GAT is retained as a control. AdamW uses learning rate 0.003, weight decay 0.0001 and batch size two. Development-parent validation selects 15, 12, 22 and seven epochs, respectively; twenty five-seed checkpoints are frozen before the new scans are extracted.',
107: 'The test uses S2, S4, S6, S7 and S8, with two opposite-corner 128³ crops per scan and voxel sizes 4.956, 8.960, 5.100, 4.803 and 4.892 μm. Eligible throats join physical pores in a topology-spanning component. Each method selects its own three highest and lowest scores plus three shared random throats. Selections are saved and hashed before hydraulic outcomes are calculated. Each throat is restricted individually at 10%, 25% and 50%; the experiment is not simultaneous multi-throat clogging.',
109: 'Fig. 4. Actual loss from individual 50% restrictions of top-selected throats. Bars show scan-balanced means; dots show the five scan means. PNM sensitivity requires a solved flow.',
110: 'Averaging first within crop and then equally across scans, guided-GAT top selections cause 3.57% mean permeability loss, versus 0.37% for random, 4.08% for affine pressure, 2.61% for guided MLP and 1.85% for guided mean GNN. GAT versus affine pressure gives exact paired sign-flip p=0.625. With five scan groups the minimum two-sided p-value is 0.0625; seed counts and intervention counts do not increase the number of independent groups.',
112: 'Independent voxel-resolved flow or measured permeability is needed to assess physical accuracy beyond the present PNM reference. This validation should use specimens from sources excluded from model development.',
114: 'Further work should test segmentation and extraction settings, voxel resolution and conductance closures under a fixed evaluation protocol. Broader source coverage is also needed to test the exploratory topology associations and evaluate response calibration or a geometry-based model selector on fresh scans.',
121: 'TopoFlow combines SNOW2 pore networks with graph attention for PNM permeability regression. Scale normalization and graph descriptors improved development performance, but physics ridge remained stronger on the locked external test: log-RMSE was 0.253, compared with 0.569 for the primary GAT.',
123: 'The throat-aware GAT identified more consequential restrictions than random selection, while affine-pressure ranking gave the larger aggregate loss. Its response predictions overestimated the reference PNM losses. These findings show that useful throat ranking does not ensure accurate response calibration.'
}
for i, text in edits.items():
    replace(i, text, 'Abstract—' if i == 4 else None)
for i in [116,118]:
    body[i].getparent().remove(body[i])
cost_heading,cost_text=edits[101].split(': ',1)
replace(101,cost_heading)
after(101,cost_text)

# Extra result paragraphs extend the existing discussion rather than replacing
# the manuscript with the separate critical-throat manuscript.
p = after(110, 'On the shared union of geometry-selected throats and three doses, 897 unique interventions yield guided-GAT response MAE 2.04 percentage points (pp), compared with 0.46 for no response, 0.75 for affine pressure and 1.92 for guided MLP. Although 78.3% of interventions cause less than 0.1 pp loss, guided GAT also underperforms for losses of at least 1 pp: MAE 6.66 versus 4.51 for no response. Ranking information therefore coexists with poor amplitude calibration. No correction is fitted on these test outcomes.')
p = after(110, 'Fresh sparse solves check 966 selected throat-dose interventions. Subsequent diagnostics reuse the same five scan groups and are explicitly post hoc. Shortest-path tortuosity has Spearman correlation 0.60 with GAT-minus-affine ranking advantage; coefficients for axial alignment, degree CV and cut density are 0.30, −0.50 and −0.30. These small-sample associations do not establish a validated topological threshold.')
# addnext inserts in reverse order; reorder the two paragraphs into a coherent
# ranking -> calibration -> post-hoc sequence.
parent = body[110].getparent()
posthoc = body[110].getnext()
calibration = posthoc.getnext()
parent.remove(calibration); body[110].addnext(calibration)
last = posthoc
for text in [
    'D. Reference and Timing Limits: Sixteen independently assembled, identical-conductance OpenPNM checks agree with the sparse reference to a maximum relative permeability difference of 2.1×10⁻¹⁴. This verifies numerical implementation. On supplied original S1/S5 maximal-ball networks, a Valvatne–Blunt shape-factor closure [14] differs from published network permeability by 12.5% and 33.5%. Closure, reservoir-boundary treatment and flow-axis differences remain material.',
    'Full-volume SNOW calculations give 342.49 and 1187.48 mD for S1 and S5, versus published image Kz values 1312 and 4440 mD. The original-network checks do not reconcile that gap. PNM target uncertainty is therefore distinct from learning error, and the restriction results remain conditional on cylindrical inscribed-diameter throats. Location-dependent clogging has prior pore-scale treatments [13]; these fixed-network interventions do not simulate evolving deposits.',
    'In the frozen restriction benchmark, one-thread warm CPU times average scan-wise median calls after five warmups and thirty repetitions. Guided GAT takes 5.31 ms per seed or 25.69 ms for a five-seed all-throat, three-dose ensemble. A base PNM solve takes 2.60 ms; three fresh 50% solves take 7.81 ms excluding the base solve. Exact all-throat precomputation takes 13.67 ms and its cached response readout 0.058 ms. SNOW2 extraction averages 6.76 s per crop. The operation scopes differ and establish no GAT speed advantage. Each restriction GAT stores 22,280 parameter bytes. Observed process RSS is 371–387 MB with all twenty models resident; this is neither peak nor incremental model memory.',
    'Named scans are the grouping units; they are not independently certified geological specimens. Four development parents, four bulk external parents and five restriction-test parents do not establish population-wide generalization. All diagnostics, numerical tests, fixed checkpoints and prediction tables are provided with the implementation at github.com/ShaunakRane914/Flow.'
]:
    if text.startswith('D.'):
        label,text=text.split(': ',1)
        e=OxmlElement('w:p');last.addnext(e);last=e
        p=Paragraph(e,doc._body);p.style='Normal';p.add_run(label)
    e = OxmlElement('w:p'); last.addnext(e); last = e
    p = Paragraph(e, doc._body); p.style = 'Normal'; p.add_run(text)
    if text.startswith('D.'):
        p.runs[0].bold = False

# Native, editable Word equations; retain KC, LBM, Reynolds and Darcy equations.
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
def el(name, *children, **attrs):
    prefix, local = name.split(':')
    node = E.Element('{' + {'w':W, 'm':M}[prefix] + '}' + local)
    for key, value in attrs.items(): node.set('{' + {'w':W, 'm':M}[prefix] + '}' + key, str(value))
    node.extend(children); return node
def r(text, normal=False):
    t = el('m:t'); t.text = text; t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    return el('m:r', el('m:rPr',el('m:sty',val='p' if normal else 'i')),
              el('w:rPr',el('w:rFonts',ascii='Cambria Math',hAnsi='Cambria Math'),el('w:sz',val='19')),t)
def seq(x): return [r(x)] if isinstance(x,str) else [x] if isinstance(x,E._Element) else x
def box(n,x): return el('m:'+n,*seq(x))
def sub(x,y): return el('m:sSub',box('e',x),box('sub',y))
def sup(x,y): return el('m:sSup',box('e',x),box('sup',y))
def ss(x,y,z): return el('m:sSubSup',box('e',x),box('sub',y),box('sup',z))
def frac(x,y): return el('m:f',box('num',x),box('den',y))
def delim(x,left='(',right=')'):
    return el('m:d',el('m:dPr',el('m:begChr',val=left),el('m:endChr',val=right)),box('e',x))
def summation(x,lo,hi=''):
    return el('m:nary',el('m:naryPr',el('m:chr',val='∑'),el('m:limLoc',val='undOvr'),
              el('m:supHide',val='0' if hi else '1')),box('sub',lo),box('sup',hi),box('e',x))
def display(index, nodes):
    p = para(index); p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    p._p.append(el('m:oMathPara',el('m:oMathParaPr',el('m:jc',val='center')),el('m:oMath',*seq(nodes))))
    return p
display(65, [r('z='),r('MLP',True),delim([frac('1',delim('V','|','|')),summation(ss('h','v','L'),'v∈V')])])
neighbors = [r('u∈N'),delim('v'),r('∪'),delim('v','{','}')]
display(67, [ss('h','v','l+1'),r('='),r('ELU',True),delim(summation([ss('α','vu','l'),sup('W','l'),ss('h','u','l')],neighbors))])
p = after(67,'')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Pt(0)
exp = lambda x: [r('exp',True),delim(x)]
p._p.append(el('m:oMathPara',el('m:oMathParaPr',el('m:jc',val='center')),el('m:oMath',sub('α','vu'),r('='),frac(exp(sub('e','vu')),summation(exp(sub('e','vj')),[r('j∈N'),delim('v'),r('∪'),delim('v','{','}')])))))
y = lambda hat=False: sub('ŷ' if hat else 'y','i')
err = lambda: sup(delim([y(),r('−'),y(True)]),'2')
display(80,[r('MSE',True),r('='),frac('1','N'),summation(err(),'i=1','N')])
display(82,[sup('R','2'),r('=1−'),frac(summation(err(),'i=1','N'),summation(sup(delim([y(),r('−ȳ')]),'2'),'i=1','N'))])
def add_equation(index,nodes):
    p=after(index,'');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
    p._p.append(el('m:oMathPara',el('m:oMathParaPr',el('m:jc',val='center')),el('m:oMath',*nodes)))
    para(index).paragraph_format.keep_with_next=True
add_equation(69,[r('ŷ=2'),sub(r('log',True),'10'),sub('d','0'),r('+'),sub('μ','t'),r('+'),sub('σ','t'),delim([r('z+'),sub('z','g')])])
add_equation(92,[r('loss',True),r('='),frac([delim('1−r'),r('s')],[r('1−'),delim('1−r'),r('ℓ')])])
for i in [22,26,64,66,74,76,79,81]: para(i).paragraph_format.keep_with_next = True

# Code-drawn scientific workflow; correct the original 256^3/GraphSAGE diagram.
im = Image.new('RGB',(1800,640),'white'); draw=ImageDraw.Draw(im)
font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf',48)
small = ImageFont.truetype('C:/Windows/Fonts/arial.ttf',36)
def arrow(a,b):
    draw.line([a,b],fill='black',width=4)
    x,y=b; x0,y0=a
    points=[(x,y),(x-16,y-10),(x-16,y+10)] if abs(x-x0)>abs(y-y0) else [(x,y),(x-10,y-16),(x+10,y-16)]
    draw.polygon(points,fill='black')
for x,text in [(20,'Segmented micro CT\n128 × 128 × 128'),(630,'SNOW2\npore network'),(1240,'Three layer GAT\nmean pooling and MLP')]:
    draw.rounded_rectangle((x,35,x+540,235),radius=10,outline='black',width=3,fill='#f5f5f5')
    draw.multiline_text((x+270,135),text,font=font,anchor='mm',align='center',spacing=14,fill='black')
arrow((565,135),(620,135));arrow((1175,135),(1230,135))
draw.rounded_rectangle((630,340,1170,540),radius=10,outline='black',width=3,fill='#f5f5f5')
draw.multiline_text((900,440),'Cylindrical throat PNM\nreference labels',font=font,anchor='mm',align='center',spacing=14,fill='black')
arrow((900,240),(900,335));arrow((1175,390),(1420,240))
draw.multiline_text((1510,440),'Predicted log10 k\n(with k in m²)',font=font,anchor='mm',align='center',spacing=14,fill='black')
arrow((1510,240),(1510,335))
draw.text((30,602),'Network extraction and model forwarding are timed separately',font=small,fill='black')
im.save(WORK/'workflow.png',dpi=(450,450))

# Compact column-width scientific plots from the committed CSVs. Native PDF
# drawing avoids shrinking presentation-sized plot labels below readable size.
def csv_rows(relative):
    with (ROOT/relative).open(newline='',encoding='utf-8') as f:
        return list(csv.DictReader(f))
def scientific_bars(name, rows, xmax, xlabel, logarithmic=False, points=None):
    width=239; height=30+19*len(rows)
    path=WORK/(name+'.pdf'); c=canvas.Canvas(str(path),pagesize=(width,height))
    left=78; right=17; bottom=27; top=height-8; span=width-left-right
    transform=(lambda v: (pymath.log10(v)+1)/5) if logarithmic else (lambda v:v/xmax)
    ticks=[0.1,1,10,100,1000,10000] if logarithmic else ([0,0.5,1,1.5,2,2.5] if xmax==2.5 else [0,5,10,15,20,25])
    c.setFont('Helvetica',7)
    for tick in ticks:
        x=left+span*transform(tick)
        c.setStrokeColor(HexColor('#dddddd'));c.setLineWidth(.35);c.line(x,bottom,x,top)
        c.setFillColor(HexColor('#000000'));c.drawCentredString(x,bottom-9,f'{tick:g}')
    for i,(label,value,sd,neural) in enumerate(rows):
        y=top-9-i*19
        c.setFillColor(HexColor('#000000'));c.setFont('Helvetica',7.5);c.drawRightString(left-4,y-2.5,label)
        c.setFillColor(HexColor('#28628b' if neural else '#999999'))
        c.rect(left,y-5,span*transform(value),10,fill=1,stroke=0)
        c.setStrokeColor(HexColor('#000000'));c.setLineWidth(.7)
        if sd:
            a=left+span*transform(max(.0001,value-sd));b=left+span*transform(value+sd)
            c.line(a,y,b,y);c.line(a,y-3,a,y+3);c.line(b,y-3,b,y+3)
        if points:
            for j,v in enumerate(points[i]):
                c.setFillColor(HexColor('#222222'));c.circle(left+span*transform(v),y+(j-2)*1.3,1.25,fill=1,stroke=0)
        if logarithmic:
            c.setFont('Helvetica',6.5);label=f'{value:.3g}'
            xpos=left+span*transform(value)+3
            if xpos+c.stringWidth(label,'Helvetica',6.5)>width-3:
                c.setFillColor(HexColor('#ffffff'));c.drawRightString(left+span*transform(value)-3,y-2,label)
            else:
                c.setFillColor(HexColor('#000000'));c.drawString(xpos,y-2,label)
    c.setStrokeColor(HexColor('#000000'));c.setLineWidth(.5);c.line(left,bottom,left,top);c.line(left,bottom,width-right,bottom)
    c.setFont('Helvetica',7.5);c.setFillColor(HexColor('#000000'));c.drawCentredString((left+width-right)/2,6,xlabel)
    c.save()
    poppler=Path(sys.executable).parents[1]/'native/poppler/Library/bin/pdftoppm.exe'
    subprocess.run([str(poppler),'-png','-singlefile','-r','450',str(path),str(WORK/name)],check=True,capture_output=True)
    return WORK/(name+'.png')

external={v['model']:v for v in csv_rows('experiments/topoflow_followup/external_test/metrics.csv')}
bulk_rows=[]
for key,label in [('physics_ridge','Physics ridge'),('pore_set_physics','Pore set'),('scaled_kc','Scaled KC'),('porosity_kc','Porosity KC'),('gat_physics','Primary GAT'),('gat_standard','Standard GAT')]:
    v=external[key];bulk_rows.append((label,float(v['rmse_log10_m2']),float(v['rmse_log10_m2_seed_std']),key.startswith('gat')))
bulk_figure=scientific_bars('external_permeability',bulk_rows,2.5,'Log10 permeability RMSE')
timings={v['stage']:float(v['median_ms']) for v in csv_rows('experiments/topoflow_research_v2/resources_verified/timing_results.csv')}
cost_rows=[(label,timings[key],0,key.endswith('forward')) for key,label in [('gat_physics_forward','GAT forward'),('sage_physics_forward','SAGE forward'),('pore_set_physics_forward','Pore set forward'),('pnm_solve','PNM solve'),('graph_conversion','Graph conversion'),('snow2','SNOW2 extraction')]]
cost_figure=scientific_bars('compute_components',cost_rows,0,'Median milliseconds on log scale',True)
rankings={v['method']:v for v in csv_rows('experiments/topoflow_discovery/discovery_metrics.csv')}
scan_ranks=csv_rows('experiments/topoflow_discovery/per_scan_discovery.csv')
rank_rows=[]; rank_points=[]
for key,label in [('guided_edge_gat','Guided GAT'),('guided_edge_sage','Guided mean GNN'),('guided_edge_mlp','Guided edge MLP'),('edge_gat','Unguided GAT'),('affine_pressure','Affine pressure'),('conductance','Conductance'),('random','Random'),('pnm_sensitivity_oracle','PNM sensitivity')]:
    rank_rows.append((label,float(rankings[key]['mean_actual_loss_pp']),0,'gat' in key))
    rank_points.append([100*float(v[key]) for v in scan_ranks])
ranking_figure=scientific_bars('critical_throat_loss',rank_rows,25,'Actual permeability loss (%)',points=rank_points)

figures = {
61: (WORK/'workflow.png',3.32),
100:(bulk_figure,3.32),
102:(cost_figure,3.32),
108:(ranking_figure,3.32)
}
for i,(path,width) in figures.items():
    p=para(i); p.clear(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(width))
after(100,'Fig. 2. Locked external permeability comparison. Error bars show five-seed standard deviation; the fitted non-neural controls are deterministic.')

def table_at(anchor, rows, widths, old=False):
    table=doc.add_table(rows=len(rows),cols=len(rows[0]))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    for col,width in zip(table.columns,widths): col.width=Inches(width)
    for ri,(row,values) in enumerate(zip(table.rows,rows)):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0: pr.append(OxmlElement('w:tblHeader'))
        for cell,value,width in zip(row.cells,values,widths):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=str(value)
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                n=OxmlElement('w:'+side);n.set(qn('w:w'),'55');n.set(qn('w:type'),'dxa');margins.append(n)
            tcpr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.keep_with_next=ri==0
                for r in p.runs:
                    r.font.name='Times New Roman';r.font.size=Pt(8);r.bold=ri==0
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom','insideH']:
        n=OxmlElement('w:'+side);n.set(qn('w:val'),'single');n.set(qn('w:sz'),'4');n.set(qn('w:color'),'808080');borders.append(n)
    props.append(borders)
    if old:
        body[anchor].addprevious(table._tbl);body[anchor].getparent().remove(body[anchor])
    else:
        body[anchor].addnext(table._tbl)
    return table

methods_table=table_at(42,[
    ['Approach','Representation','Reference and scope'],
    ['KC [1]','Bulk descriptors','Analytical family; calibrated variants evaluated'],
    ['LBM','Voxel flow','Direct simulation; not run in this study'],
    ['CNN [6]','Voxel features','Prior image regression; no matched CNN benchmark'],
    ['GNN [7]–[10]','Graph representations','CNN feature graphs or pore–throat graphs; transfer and restrictions tested here']
],[0.65,0.88,1.78],old=True)
features_table=table_at(71,[
    ['Feature group','Bulk GAT','Restriction GAT'],
    ['Pore geometry','Log diameter and volume','Normalized diameter and volume'],
    ['Position and boundary','Not used','3 coordinates and 2 boundary flags'],
    ['Throat geometry','Not used','Log diameter, length and d⁴/L proxy'],
    ['Throat direction','Not used','3 absolute direction components'],
    ['Global features','10 descriptors in ablation','Global edge-score normalization']
],[0.99,1.10,1.22],old=True)
for index,table in [(43,methods_table),(72,features_table)]:
    body[index].getparent().remove(body[index]);table._tbl.addprevious(body[index])
    para(index).paragraph_format.keep_with_next=True

with (ROOT/'experiments/topoflow_followup/external_test/metrics.csv').open(newline='',encoding='utf-8') as f:
    external={r['model']:r for r in csv.DictReader(f)}
metric_rows=[['Model','Log RMSE','Log R²','Within ×2']]
for key,label in [('physics_ridge','Physics ridge'),('pore_set_physics','Pore set + geometry'),('scaled_kc','Scaled KC'),('gat_descriptors_only','GAT descriptors'),('gat_normalized_only','GAT normalized'),('porosity_kc','Porosity KC'),('gat_physics','Primary GAT'),('gat_standard','Standard GAT')]:
    v=external[key];std=float(v['rmse_log10_m2_seed_std'])
    rmse=f"{float(v['rmse_log10_m2']):.3f}"
    if std: rmse+=f" ± {std:.3f}"
    metric_rows.append([label,rmse,f"{float(v['r2_log10_m2']):.3f}",f"{100*float(v['factor_of_two_accuracy']):.1f}%"])
table_at(104,metric_rows,[1.12,1.03,0.56,0.60])

# Replace only the old bibliography, retaining its location and Word appearance.
references = [
'[1] R. P. Chapuis and M. Aubertin, “On the use of the Kozeny–Carman equation to predict the hydraulic conductivity of soils,” Canadian Geotechnical Journal, vol. 40, no. 3, pp. 616–628, 2003, doi: 10.1139/t03-013.',
'[2] H. Dong and M. J. Blunt, “Pore-network extraction from micro-computerized-tomography images,” Physical Review E, vol. 80, 036307, 2009, doi: 10.1103/PhysRevE.80.036307.',
'[3] J. T. Gostick, “Versatile and efficient pore network extraction method using marker-based watershed segmentation,” Physical Review E, vol. 96, 023307, 2017, doi: 10.1103/PhysRevE.96.023307.',
'[4] J. Gostick et al., “OpenPNM: A pore network modeling package,” Computing in Science & Engineering, vol. 18, no. 4, pp. 60–74, 2016, doi: 10.1109/MCSE.2016.49.',
'[5] P. Veličković et al., “Graph Attention Networks,” International Conference on Learning Representations, 2018, arXiv:1710.10903.',
'[6] J. Wu, X. Yin and H. Xiao, “Seeing permeability from images: Fast prediction with convolutional neural networks,” Science Bulletin, vol. 63, no. 18, pp. 1215–1222, 2018, doi: 10.1016/j.scib.2018.08.006.',
'[7] M. K. Alzahrani, A. Shapoval, Z. Chen et al., “Pore-GNN: A graph neural network-based framework for predicting flow properties of porous media from micro-CT images,” Advances in Geo-Energy Research, vol. 10, no. 1, pp. 39–55, 2023, doi: 10.46690/ager.2023.10.05.',
'[8] X. Zhao, Y. Zhong and P. Li, “RTG-GNN: A novel rock topology-guided approach for permeability prediction using graph neural networks,” Geoenergy Science and Engineering, vol. 243, 213358, 2024, doi: 10.1016/j.geoen.2024.213358.',
'[9] J. M. I’Anson, M. J. H. Simmons, E. H. Stitt and R. W. Gallen, “Using graph neural networks to predict the permeability of porous media,” InterPore Journal, vol. 2, no. 3, 2025, doi: 10.69631/ipj.v2i3nr47.',
'[10] Q. Zhao and H. Xiao, “An end-to-end differentiable, graph neural network-embedded pore network model for permeability prediction,” Journal of Hydrology, vol. 675, 135602, 2026, doi: 10.1016/j.jhydrol.2026.135602.',
'[11] Imperial College London, Pore-scale Modelling Group, “ImagesICPSC2007,” binary images and source headers. Accessed Oct. 22, 2026. Available: https://imperialcollegelondon.box.com/v/ImagesICPSC2007.',
'[12] Imperial College London, “Estaillades 1000³ micro-CT image, 3.31136 μm.” Accessed Oct. 22, 2026. Official archive: https://imperialcollegelondon.box.com/shared/static/2n4lb95amz7fstk19q826uzu922eb4b1.zip.',
'[13] S. Nassiri, L. van Paassen and N. Mahabadi, “Evolution of porous media and porosity–permeability relations during biogeochemical clogging processes: A pore-scale simulation study,” Computers and Geotechnics, vol. 198, 108292, 2026, doi: 10.1016/j.compgeo.2026.108292.',
'[14] OpenPNM, “Valvatne–Blunt hydraulic conductance,” official implementation documentation, installed version 3.6.0. Available: https://openpnm.org/_modules/openpnm/models/physics/hydraulic_conductance/_funcs.html#valvatne_blunt. Accessed Oct. 22, 2026.',
'[15] W. L. Hamilton, R. Ying and J. Leskovec, “Inductive representation learning on large graphs,” Advances in Neural Information Processing Systems, vol. 30, 2017, arXiv:1706.02216.',
'[16] Y. H. Qian, D. d’Humières and P. Lallemand, “Lattice BGK models for Navier–Stokes equation,” Europhysics Letters, vol. 17, no. 6, pp. 479–484, 1992, doi: 10.1209/0295-5075/17/6/001.',
'[17] J. Sherman and W. J. Morrison, “Adjustment of an inverse matrix corresponding to a change in one element of a given matrix,” The Annals of Mathematical Statistics, vol. 21, no. 1, pp. 124–127, 1950, doi: 10.1214/aoms/1177729893.'
]
for idx,text in enumerate(references,125):
    p=replace(idx,text)
    p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.space_after=Pt(4)
    for r in p.runs: r.font.size=Pt(8)
for idx in range(125+len(references),150):
    body[idx].getparent().remove(body[idx])

# Remove empty spacer paragraphs and duplicate figure/table lists; retain every
# section break and all original author and page-geometry elements.
for idx in range(150,164):
    e=body[idx]
    if e.getparent() is not None and not e.xpath('.//w:sectPr'):
        e.getparent().remove(e)
for idx,e in enumerate(body):
    if e.getparent() is None or e.tag!=qn('w:p'): continue
    p=Paragraph(e,doc._body)
    if not p.text.strip() and not e.xpath('.//w:drawing | .//m:oMath | .//w:sectPr'):
        e.getparent().remove(e)

# Keep equations and figures with their introductions/captions, but let body
# paragraphs flow naturally through the source's two-column sections.
for idx in [43,62,72,103,109]:
    p=para(idx);p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.space_after=Pt(5)
    for r in p.runs:r.font.size=Pt(8)
para(104).paragraph_format.keep_with_next=True

# The source mixes 6-point runs and whole-paragraph italics. Retain its layout
# while regularizing body typography and suppressing duplicated auto letters.
major={8,19,45,51,93,111,120,124}
subheads={52,56,84,94,101,105}
captions={43,62,72,103,104,109}
inline={20,25,29,30,63,73}
equations={23,27,65,67,75,77,80,82}
element_indices={id(e):i for i,e in enumerate(body)}
for p in doc.paragraphs:
    index=element_indices.get(id(p._p))
    if index in major or index==0 or p._p.xpath('.//w:sectPr'):
        continue
    if index in range(125,125+len(references)):
        p.style='Normal';p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent=Pt(12);p.paragraph_format.first_line_indent=Pt(-12)
        p.paragraph_format.line_spacing=1.0
        p.paragraph_format.keep_together=True
        for r in p.runs:r.italic=False;r.font.size=Pt(8)
        continue
    if index in {4,5,6,7}:
        p.paragraph_format.first_line_indent=Pt(0)
        for r in p.runs:r.font.size=Pt(9)
        continue
    if p._p.xpath('.//w:drawing | .//m:oMath'):
        continue
    pr=p._p.get_or_add_pPr()
    for n in pr.findall(qn('w:numPr')):pr.remove(n)
    num=OxmlElement('w:numPr');ident=OxmlElement('w:numId');ident.set(qn('w:val'),'0');num.append(ident);pr.append(num)
    p.style='Normal'
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing=1.0
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.first_line_indent=Pt(10)
    for r in p.runs:
        r.font.name='Times New Roman';r.font.size=Pt(10);r.italic=False;r.bold=False
    if index in captions or p.text.startswith('Fig. 2.'):
        p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:r.font.size=Pt(8)
    elif index in subheads or p.text=='D. Reference and Timing Limits':
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.space_before=Pt(6)
        p.paragraph_format.keep_with_next=True
        for r in p.runs:r.italic=True
    elif index in inline:
        p.paragraph_format.first_line_indent=Pt(0)
        text=p.text;prefix,rest=text.split(':',1);p.clear()
        a=p.add_run(prefix+':');a.italic=True
        p.add_run(rest)
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10)

# Number references in their first-citation order, including table citations.
citation_order=[]
for child in doc._element.body:
    if child is body[124]:break
    for node in child.iter(qn('w:t')):
        for match in re.finditer(r'\[(\d+)\]',node.text or ''):
            old=int(match.group(1))
            if old not in citation_order:citation_order.append(old)
assert set(citation_order)==set(range(1,len(references)+1))
citation_map={old:i for i,old in enumerate(citation_order,1)}
for node in doc._element.body.iter(qn('w:t')):
    if node.text:
        node.text=re.sub(r'\[(\d+)\]',lambda m:'['+str(citation_map[int(m.group(1))])+']',node.text)
last=body[124]
for old in citation_order:
    element=body[124+old]
    element.getparent().remove(element);last.addnext(element);last=element
doc.core_properties.title=title.text
doc.core_properties.author='Shaunak Rane; Tanisha Somai; Hemant Maddu'
doc.core_properties.subject='Source-separated evaluation of TopoFlow permeability prediction'
doc.core_properties.comments='Existing manuscript revised using completed, committed experiments.'
assert E.tostring(body[1])==original_author, 'Author block was modified'
assert [E.tostring(e) for e in doc._element.xpath('//w:sectPr')]==sections_before, 'Page sections changed'
assert len(doc._element.xpath('//m:oMathPara'))==11
assert len(doc.inline_shapes)==4
# Remove the old, now-unreferenced illustration parts from the edited package.
used_images={n.get(qn('r:embed')) for n in doc._element.xpath('//*[@r:embed]')}
for rid,rel in list(doc.part.rels.items()):
    if rel.reltype.endswith('/image') and rid not in used_images:
        doc.part.drop_rel(rid)
# A closing continuous section balances the final two-column page. All four
# retained section properties remain unchanged; the new empty tail has one column.
final_section=doc._element.body.sectPr
closing=OxmlElement('w:p');pp=OxmlElement('w:pPr')
pp.append(deepcopy(final_section));closing.append(pp)
doc._element.body.insert(len(doc._element.body)-1,closing)
final_section.find(qn('w:cols')).set(qn('w:num'),'1')
doc.save(OUTPUT)
with ZipFile(OUTPUT) as z: assert z.testzip() is None
proof={
    'source':SOURCE.name,'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'output':OUTPUT.name,'output_sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
    'changed_paragraphs':len(changes),'preserved_author_block':True,
    'preserved_section_properties':True,'native_display_equations':11,'figures':4,
    'closing_balancing_section':True,'reference_citation_map':citation_map,
    'evidence_directories':['experiments/topoflow_research_v2','experiments/topoflow_followup','experiments/topoflow_discovery','experiments/topoflow_diagnostics']
}
(WORK/'revision_audit.json').write_text(json.dumps(proof,indent=2),encoding='utf-8')
(WORK/'paragraph_changes.json').write_text(json.dumps(changes,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps(proof,indent=2))
